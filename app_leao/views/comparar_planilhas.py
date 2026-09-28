import io
import re
import unicodedata
from datetime import datetime, timedelta

from django.contrib import messages
from django.http import HttpResponse
from django.shortcuts import render
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd


# ==============================================================================
# DICIONÁRIO / DEPARA DE CONTAS CORRENTES
# ==============================================================================
MAPEAMENTO_CONTA_CORRENTE = {
    "40147": "Bradesco (Matriz)",
    "40436": "Bradesco (Aeroporto)",
    "41835": "Bradesco (Castanheira)",
    "41828": "Bradesco (Baenão)",
    "41818": "Bradesco (Pátio Belém)",
}

DEPARA_REVERSO = {
    v.upper(): v for k, v in MAPEAMENTO_CONTA_CORRENTE.items()
}
DEPARA_REVERSO.update(
    {
        "BRADESCO (AERO)": "Bradesco (Aeroporto)",
        "BRADESCO (E-COMMERCE)": "Bradesco (Matriz)",
        "LEÃO AZUL WEBSTORE LTDA": "Bradesco (Matriz)",
    }
)


def _formatar_conta_corrente(valor_bruto) -> str:
    """Padroniza códigos ou nomes de contas correntes conforme a regra definida."""
    if pd.isna(valor_bruto) or not valor_bruto:
        return "Não Identificado"

    val_str = str(valor_bruto).strip().replace(".0", "")

    # 1. Busca direta por código da conta
    for cod, nome in MAPEAMENTO_CONTA_CORRENTE.items():
        if cod in val_str:
            return nome

    # 2. Busca reverso por nome/apelido
    val_norm = val_str.upper()
    for chave_ref, nome_correto in DEPARA_REVERSO.items():
        if chave_ref in val_norm:
            return nome_correto

    return val_str


# ==============================================================================
# FUNÇÕES DE HIGIENIZAÇÃO DE DADOS
# ==============================================================================
def _normalizar_texto(texto: str) -> str:
    if not texto or pd.isna(texto):
        return ""
    txt = str(texto).upper()
    txt = (
        unicodedata.normalize("NFKD", txt)
        .encode("ascii", "ignore")
        .decode("ascii")
    )
    txt = re.sub(r"[^\w\s]", " ", txt)
    return " ".join(txt.split())


def _limpar_valor_monetario(val) -> float:
    if pd.isna(val) or val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return round(abs(float(val)), 2)

    s = (
        str(val)
        .strip()
        .replace("]", "")
        .replace("R$", "")
        .replace(" ", "")
    )
    if not s:
        return 0.0

    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")

    try:
        return round(abs(float(s)), 2)
    except ValueError:
        return 0.0


def _formatar_linha_digitavel(val) -> str:
    """Garante que a linha digitável seja mantida como string sem notação científica."""
    if pd.isna(val) or val is None:
        return ""
    val_str = str(val).strip()

    if "E+" in val_str.upper():
        try:
            val_str = f"{float(val_str):.0f}"
        except ValueError:
            pass

    if val_str.endswith(".0"):
        val_str = val_str[:-2]

    return val_str


# ==============================================================================
# PREPARAÇÃO DA PLANILHA DE PROVISÃO
# ==============================================================================
def _preparar_dataframe_provisao(arquivo_django) -> pd.DataFrame:
    if hasattr(arquivo_django, "seek"):
        arquivo_django.seek(0)

    df = pd.read_excel(arquivo_django, dtype=str)

    col_venc = next(
        (
            c
            for c in df.columns
            if "VENC" in _normalizar_texto(str(c))
            or "DATA" in _normalizar_texto(str(c))
        ),
        df.columns[0],
    )
    col_val = next(
        (
            c
            for c in df.columns
            if "VALOR" in _normalizar_texto(str(c))
            or "TOTAL" in _normalizar_texto(str(c))
        ),
        df.columns[1],
    )
    col_forn = next(
        (
            c
            for c in df.columns
            if "FORNECEDOR" in _normalizar_texto(str(c))
            or "DESCRICAO" in _normalizar_texto(str(c))
        ),
        df.columns[2],
    )
    col_loja = next(
        (
            c
            for c in df.columns
            if "LOJA" in _normalizar_texto(str(c))
            or "CONTA" in _normalizar_texto(str(c))
            or "PAGADOR" in _normalizar_texto(str(c))
        ),
        None,
    )

    df = df[df[col_venc].astype(str).str.upper() != "TOTAL"].copy()

    df["Data_Norm"] = pd.to_datetime(
        df[col_venc], dayfirst=True, errors="coerce"
    ).dt.date
    df["Valor_Norm"] = df[col_val].apply(_limpar_valor_monetario)
    df["Fornecedor_Norm"] = df[col_forn].apply(_normalizar_texto)

    if col_loja:
        df["Conta_Formatada"] = df[col_loja].apply(_formatar_conta_corrente)
    else:
        df["Conta_Formatada"] = "Não Identificado"

    res = df[df["Valor_Norm"] > 0].copy()
    res.attrs["col_forn"] = col_forn
    return res


# ==============================================================================
# BIFURCAÇÃO 1: MODO BOLETOS (BOLETOS EM MÃOS vs PROVISÃO)
# ==============================================================================
def conciliar_boletos_em_maos(arq1, arq2) -> io.BytesIO:
    if hasattr(arq1, "seek"):
        arq1.seek(0)
    if hasattr(arq2, "seek"):
        arq2.seek(0)

    df_prov = _preparar_dataframe_provisao(arq1)
    df_boletos_orig = pd.read_excel(arq2, dtype=str)

    col_venc_b = next(
        (
            c
            for c in df_boletos_orig.columns
            if "VENC" in _normalizar_texto(str(c))
            or "DATA" in _normalizar_texto(str(c))
        ),
        df_boletos_orig.columns[0],
    )
    col_val_b = next(
        (
            c
            for c in df_boletos_orig.columns
            if "VALOR" in _normalizar_texto(str(c))
        ),
        df_boletos_orig.columns[1],
    )
    col_pagador_b = next(
        (
            c
            for c in df_boletos_orig.columns
            if "PAGADOR" in _normalizar_texto(str(c))
            or "CONTA" in _normalizar_texto(str(c))
            or "LOJA" in _normalizar_texto(str(c))
        ),
        None,
    )

    df_boletos_orig = df_boletos_orig[
        df_boletos_orig[col_venc_b].astype(str).str.upper() != "TOTAL"
    ].copy()

    if col_pagador_b:
        df_boletos_orig[col_pagador_b] = df_boletos_orig[col_pagador_b].apply(
            _formatar_conta_corrente
        )

    df_boletos_work = df_boletos_orig.copy()
    df_boletos_work["_Data_Norm"] = pd.to_datetime(
        df_boletos_work[col_venc_b], dayfirst=True, errors="coerce"
    ).dt.date
    df_boletos_work["_Valor_Norm"] = df_boletos_work[col_val_b].apply(
        _limpar_valor_monetario
    )

    df_boletos_work = df_boletos_work[
        df_boletos_work["_Valor_Norm"] > 0
    ].copy()

    # LÓGICA INVERTIDA: A menor data de referência é a MENOR DATA DA PLANILHA DE BOLETOS
    data_min_boletos = df_boletos_work["_Data_Norm"].min()

    df_prov["_utilizado"] = False
    df_boletos_work["_status_match"] = "NÃO PROVISIONADO"
    df_boletos_work["_observacao"] = ""

    for idx, b in df_boletos_work.iterrows():
        dt = b["_Data_Norm"]
        vl = b["_Valor_Norm"]

        # Se a planilha de provisão tiver datas abaixo de data_min_boletos, a busca vai ignorar registros antigos da provisão
        # 1. Match Exato (considerando provisão na janela válida)
        match1 = df_prov[
            (~df_prov["_utilizado"])
            & (df_prov["Data_Norm"] >= data_min_boletos)
            & (df_prov["Data_Norm"] == dt)
            & (df_prov["Valor_Norm"] == vl)
        ]
        if not match1.empty:
            m_idx = match1.index[0]
            df_prov.loc[m_idx, "_utilizado"] = True
            df_boletos_work.loc[idx, "_status_match"] = "CONFERIDO"
            df_boletos_work.loc[idx, "_observacao"] = (
                "Match 100% Exato de Valor e Vencimento"
            )
            continue

        # 2. Tolerância de Vencimento (±3 dias)
        if dt is not None and not pd.isna(dt):
            dt_min = dt - timedelta(days=3)
            dt_max = dt + timedelta(days=3)
            match2 = df_prov[
                (~df_prov["_utilizado"])
                & (df_prov["Data_Norm"] >= data_min_boletos)
                & (df_prov["Data_Norm"] >= dt_min)
                & (df_prov["Data_Norm"] <= dt_max)
                & (df_prov["Valor_Norm"] == vl)
            ]
            if not match2.empty:
                m_idx = match2.index[0]
                df_prov.loc[m_idx, "_utilizado"] = True
                df_boletos_work.loc[idx, "_status_match"] = "CONFERIDO"
                dt_p = df_prov.loc[m_idx, "Data_Norm"]
                df_boletos_work.loc[idx, "_observacao"] = (
                    f"Valor Exato (Vencimento Provisão: {dt_p.strftime('%d/%m/%Y')})"
                )
                continue

        # 3. Tolerância de Valor (±R$ 5,00)
        match3 = df_prov[
            (~df_prov["_utilizado"])
            & (df_prov["Data_Norm"] >= data_min_boletos)
            & (df_prov["Data_Norm"] == dt)
            & (df_prov["Valor_Norm"] >= vl - 5.0)
            & (df_prov["Valor_Norm"] <= vl + 5.0)
        ]
        if not match3.empty:
            m_idx = match3.index[0]
            df_prov.loc[m_idx, "_utilizado"] = True
            df_boletos_work.loc[idx, "_status_match"] = "INCONCLUINTE"
            v_p = df_prov.loc[m_idx, "Valor_Norm"]
            df_boletos_work.loc[idx, "_observacao"] = (
                f"Vencimento igual, mas valor na Provisão é R$ {v_p:.2f}"
            )
            continue

    idx_conf = df_boletos_work[
        df_boletos_work["_status_match"] == "CONFERIDO"
    ].index
    idx_inconc = df_boletos_work[
        df_boletos_work["_status_match"] == "INCONCLUINTE"
    ].index
    idx_np = df_boletos_work[
        df_boletos_work["_status_match"] == "NÃO PROVISIONADO"
    ].index

    # Identifica registros da Provisão que ficaram abaixo da menor data da planilha de boletos
    df_prov_fora_periodo = df_prov[
        (df_prov["Data_Norm"] < data_min_boletos)
    ].copy()
    if not df_prov_fora_periodo.empty and data_min_boletos:
        df_prov_fora_periodo["Motivo"] = (
            f"Registro na Provisão anterior à data inicial dos boletos ({data_min_boletos.strftime('%d/%m/%Y')})"
        )

    df_conferidos = df_boletos_orig.loc[idx_conf].copy()
    df_conferidos["Observação do Match"] = df_boletos_work.loc[
        idx_conf, "_observacao"
    ]

    df_inconcluintes = df_boletos_orig.loc[idx_inconc].copy()
    df_inconcluintes["Análise de Dúvida"] = df_boletos_work.loc[
        idx_inconc, "_observacao"
    ]

    df_nao_prov = df_boletos_orig.loc[idx_np].copy()

    wb = openpyxl.Workbook()

    cor_azul = "1A365D"
    cor_verde = "2F855A"
    cor_amarelo = "D69E2E"
    cor_vermelho = "C53030"
    cor_cinza_escuro = "4A5568"
    cinza_fundo = "EDF2F7"
    cinza_borda = "CBD5E0"

    fonte_bold = Font(name="Segoe UI", size=10, bold=True)
    fonte_base = Font(name="Segoe UI", size=10)
    borda_fina = Border(
        left=Side(border_style="thin", color=cinza_borda),
        right=Side(border_style="thin", color=cinza_borda),
        top=Side(border_style="thin", color=cinza_borda),
        bottom=Side(border_style="thin", color=cinza_borda),
    )

    ws_resumo = wb.active
    ws_resumo.title = "Resumo da Conciliação"
    ws_resumo.views.sheetView[0].showGridLines = True

    ws_resumo.merge_cells("A1:E1")
    b = ws_resumo["A1"]
    b.value = "PAINEL DE AUDITORIA DE BOLETOS"
    b.font = Font(name="Segoe UI", size=12, bold=True, color="FFFFFF")
    b.fill = PatternFill(
        start_color=cor_azul, end_color=cor_azul, fill_type="solid"
    )
    b.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_resumo.row_dimensions[1].height = 28

    headers_resumo = [
        "SITUAÇÃO DO BOLETO",
        "QTD REGISTROS",
        "VALOR TOTAL (R$)",
        "% DO TOTAL",
    ]
    ws_resumo.row_dimensions[3].height = 20
    for c_i, h in enumerate(headers_resumo, start=1):
        cell = ws_resumo.cell(row=3, column=c_i, value=h)
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            start_color="2B6CB0", end_color="2B6CB0", fill_type="solid"
        )
        cell.alignment = Alignment(horizontal="center", vertical="center")

    tot_qtd = len(df_boletos_work)
    linhas_resumo = [
        (
            "BOLETOS NÃO PROVISIONADOS (PARA IMPORTAR NO SISTEMA)",
            len(df_nao_prov),
            float(df_boletos_work.loc[idx_np, "_Valor_Norm"].sum()),
            cor_vermelho,
        ),
        (
            "INCONCLUINTE / VERIFICAR MANUAMENTE",
            len(df_inconcluintes),
            float(df_boletos_work.loc[idx_inconc, "_Valor_Norm"].sum()),
            cor_amarelo,
        ),
        (
            "CONFERIDOS E PROVISIONADOS",
            len(df_conferidos),
            float(df_boletos_work.loc[idx_conf, "_Valor_Norm"].sum()),
            cor_verde,
        ),
        (
            "PROVISÃO ANTERIOR À DATA DOS BOLETOS (FORA DO PERÍODO)",
            len(df_prov_fora_periodo),
            float(df_prov_fora_periodo["Valor_Norm"].sum()) if not df_prov_fora_periodo.empty else 0.0,
            cor_cinza_escuro,
        ),
    ]

    for idx_row, (st, q, v, cor) in enumerate(linhas_resumo, start=4):
        ws_resumo.row_dimensions[idx_row].height = 20
        c1 = ws_resumo.cell(row=idx_row, column=1, value=st)
        c1.font = Font(name="Segoe UI", size=10, bold=True, color=cor)

        c2 = ws_resumo.cell(row=idx_row, column=2, value=q)
        c2.font = fonte_base
        c2.alignment = Alignment(horizontal="center")

        c3 = ws_resumo.cell(row=idx_row, column=3, value=v)
        c3.font = fonte_base
        c3.number_format = "R$ #,##0.00;[Red]-R$ #,##0.00"

        pct = (q / tot_qtd) if tot_qtd > 0 else 0
        c4 = ws_resumo.cell(row=idx_row, column=4, value=pct)
        c4.font = fonte_base
        c4.number_format = "0.0%"
        c4.alignment = Alignment(horizontal="center")

        for c in [c1, c2, c3, c4]:
            c.border = borda_fina

    r_tot = 8
    ws_resumo.row_dimensions[r_tot].height = 22
    ws_resumo.cell(row=r_tot, column=1, value="TOTAL BOLETOS RECEBIDOS").font = (
        fonte_bold
    )
    ws_resumo.cell(row=r_tot, column=2, value="=SUM(B4:B6)").font = fonte_bold
    ws_resumo.cell(row=r_tot, column=3, value="=SUM(C4:C6)").font = fonte_bold
    ws_resumo.cell(row=r_tot, column=3).number_format = (
        "R$ #,##0.00;[Red]-R$ #,##0.00"
    )
    ws_resumo.cell(row=r_tot, column=4, value="=SUM(D4:D6)").font = fonte_bold
    ws_resumo.cell(row=r_tot, column=4).number_format = "0.0%"

    for col_i in range(1, 5):
        c = ws_resumo.cell(row=r_tot, column=col_i)
        c.fill = PatternFill(
            start_color=cinza_fundo, end_color=cinza_fundo, fill_type="solid"
        )
        c.border = Border(
            top=Side(border_style="thin", color=cinza_borda),
            bottom=Side(border_style="double", color="1A202C"),
        )

    def _exportar_tabela_original(
        nome_aba, df_dados, cor_banner, subtitulo_banner
    ):
        ws = wb.create_sheet(title=nome_aba)
        ws.views.sheetView[0].showGridLines = True

        cols_orig = list(df_dados.columns)
        num_cols = len(cols_orig)
        letra_max = get_column_letter(max(num_cols, 1))

        ws.merge_cells(f"A1:{letra_max}1")
        banner = ws["A1"]
        banner.value = f"{subtitulo_banner.upper()}"
        banner.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        banner.fill = PatternFill(
            start_color=cor_banner, end_color=cor_banner, fill_type="solid"
        )
        banner.alignment = Alignment(
            horizontal="left", vertical="center", indent=1
        )
        ws.row_dimensions[1].height = 24

        ws.row_dimensions[3].height = 20
        for c_idx, h in enumerate(cols_orig, start=1):
            cell = ws.cell(row=3, column=c_idx, value=str(h))
            cell.font = Font(
                name="Segoe UI", size=10, bold=True, color="FFFFFF"
            )
            cell.fill = PatternFill(
                start_color="4A5568", end_color="4A5568", fill_type="solid"
            )
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for i, row in df_dados.reset_index(drop=True).iterrows():
            l_atual = 4 + i
            ws.row_dimensions[l_atual].height = 19
            bg_zebra = PatternFill(
                start_color="F8FAFC" if i % 2 == 1 else "FFFFFF",
                fill_type="solid",
            )

            for c_idx, col_name in enumerate(cols_orig, start=1):
                val_raw = row[col_name]
                c = ws.cell(row=l_atual, column=c_idx)
                col_upper = str(col_name).upper()

                es_coluna_valor = any(
                    kw in col_upper for kw in ["VALOR", "TOTAL", "MONTANTE", "PRECO", "SALDO", "LIQUIDO"]
                )

                if es_coluna_valor:
                    val_num = _limpar_valor_monetario(val_raw)
                    c.value = val_num
                    c.number_format = "R$ #,##0.00;[Red]-R$ #,##0.00"
                    c.alignment = Alignment(horizontal="right")
                elif any(kw in col_upper for kw in ["BOLETO", "LINHA", "CODIGO", "CPF", "CNPJ", "BARRA"]):
                    val_limpo = _formatar_linha_digitavel(val_raw)
                    c.value = val_limpo
                    c.number_format = "@"
                    c.alignment = Alignment(horizontal="center")
                elif any(kw in col_upper for kw in ["VENC", "DATA"]):
                    c.value = str(val_raw) if pd.notna(val_raw) and val_raw != "" else ""
                    c.alignment = Alignment(horizontal="center")
                else:
                    if isinstance(val_raw, (int, float)) and not pd.isna(val_raw):
                        c.value = val_raw
                    else:
                        c.value = str(val_raw) if pd.notna(val_raw) and val_raw != "" else ""
                        c.alignment = Alignment(horizontal="left")

                c.border = borda_fina
                c.font = fonte_base
                c.fill = bg_zebra

        if not df_dados.empty:
            l_tot = 4 + len(df_dados)
            ws.row_dimensions[l_tot].height = 22
            ws.cell(row=l_tot, column=1, value="TOTAL").font = fonte_bold

            for c_idx, col_name in enumerate(cols_orig, start=1):
                c = ws.cell(row=l_tot, column=c_idx)
                c.fill = PatternFill(
                    start_color=cinza_fundo,
                    end_color=cinza_fundo,
                    fill_type="solid",
                )
                c.border = Border(
                    top=Side(border_style="thin", color=cinza_borda),
                    bottom=Side(border_style="double", color="1A202C"),
                )

                col_upper = str(col_name).upper()
                if any(kw in col_upper for kw in ["VALOR", "TOTAL", "MONTANTE", "PRECO", "SALDO", "LIQUIDO"]):
                    letra_col = get_column_letter(c_idx)
                    c.value = f"=SUM({letra_col}4:{letra_col}{l_tot-1})"
                    c.font = fonte_bold
                    c.number_format = "R$ #,##0.00;[Red]-R$ #,##0.00"
                    c.alignment = Alignment(horizontal="right")

        for col_idx in range(1, num_cols + 1):
            comp_max = max(
                len(str(ws.cell(row=r, column=col_idx).value or ""))
                for r in range(3, ws.max_row + 1)
            )
            ws.column_dimensions[get_column_letter(col_idx)].width = max(
                comp_max + 4, 15
            )

    _exportar_tabela_original(
        "Boletos a Importar",
        df_nao_prov,
        cor_vermelho,
        "BOLETOS NÃO ENCONTRADOS NO PERÍODO (PRONTOS PARA IMPORTAR NO SISTEMA)",
    )
    _exportar_tabela_original(
        "Inconcluintes",
        df_inconcluintes,
        cor_amarelo,
        "BOLETOS COM DIVERGÊNCIA LEVE DE VALOR (VERIFICAR MANUAMENTE)",
    )
    _exportar_tabela_original(
        "Boletos Conferidos",
        df_conferidos,
        cor_verde,
        "BOLETOS CONFERIDOS COM SUCESSO NO PROVISIONADO",
    )
    _exportar_tabela_original(
        "Fora do Período",
        df_prov_fora_periodo,
        cor_cinza_escuro,
        "REGISTROS DE PROVISÃO DE MESES ANTERIORES À MENOR DATA DA PLANILHA DE BOLETOS",
    )

    for col_idx in range(1, 5):
        ws_resumo.column_dimensions[get_column_letter(col_idx)].width = 45

    output_stream = io.BytesIO()
    wb.save(output_stream)
    output_stream.seek(0)
    return output_stream


# ==============================================================================
# BIFURCAÇÃO 2: MODO BANCO (EXTRATO BANCÁRIO vs PROVISÃO)
# ==============================================================================
def conciliar_extrato_banco(arq1, arq2) -> io.BytesIO:
    df1 = _preparar_dataframe_provisao(arq1)

    if hasattr(arq2, "seek"):
        arq2.seek(0)
    df2 = pd.read_excel(arq2, dtype=str)

    col_venc_2 = next(
        (
            c
            for c in df2.columns
            if "VENC" in _normalizar_texto(str(c))
            or "DATA" in _normalizar_texto(str(c))
        ),
        df2.columns[0],
    )
    col_val_2 = next(
        (
            c
            for c in df2.columns
            if "VALOR" in _normalizar_texto(str(c))
        ),
        df2.columns[1],
    )
    col_forn_2 = next(
        (
            c
            for c in df2.columns
            if "FORNECEDOR" in _normalizar_texto(str(c))
            or "HISTORICO" in _normalizar_texto(str(c))
            or "DESCRICAO" in _normalizar_texto(str(c))
        ),
        df2.columns[2],
    )
    col_conta_2 = next(
        (
            c
            for c in df2.columns
            if "CONTA" in _normalizar_texto(str(c))
            or "CC" in _normalizar_texto(str(c))
            or "PAGADOR" in _normalizar_texto(str(c))
        ),
        None,
    )

    df2["Data_Norm"] = pd.to_datetime(
        df2[col_venc_2], dayfirst=True, errors="coerce"
    ).dt.date
    df2["Valor_Norm"] = df2[col_val_2].apply(_limpar_valor_monetario)
    df2["Fornecedor_Norm"] = df2[col_forn_2].apply(_normalizar_texto)

    if col_conta_2:
        df2["Conta_Norm"] = df2[col_conta_2].apply(_formatar_conta_corrente)
    else:
        df2["Conta_Norm"] = "Não Identificado"

    df2["_utilizado"] = False
    status_lista = []
    obs_lista = []
    conta_paga_lista = []
    data_paga_lista = []

    for idx1, row1 in df1.iterrows():
        dt1 = row1["Data_Norm"]
        vl1 = row1["Valor_Norm"]

        if pd.isna(dt1) or vl1 == 0:
            status_lista.append("DADOS_INCOMPLETOS")
            obs_lista.append("Data inválida ou valor zerado")
            conta_paga_lista.append("-")
            data_paga_lista.append(None)
            continue

        dt_min = dt1 - timedelta(days=35)
        dt_max = dt1 + timedelta(days=30)

        candidatos_p2 = df2[
            (~df2["_utilizado"])
            & (df2["Valor_Norm"] == vl1)
            & (df2["Data_Norm"] >= dt_min)
            & (df2["Data_Norm"] <= dt_max)
        ].copy()

        if not candidatos_p2.empty:
            candidatos_p2["_dias_diff"] = candidatos_p2["Data_Norm"].apply(
                lambda d: abs((d - dt1).days)
            )
            candidatos_p2 = candidatos_p2.sort_values(
                by=["_dias_diff"], ascending=True
            )

            idx_m = candidatos_p2.index[0]
            match_row = candidatos_p2.loc[idx_m]
            df2.loc[idx_m, "_utilizado"] = True

            dias_diff = (match_row["Data_Norm"] - dt1).days
            status_lista.append("EXATO")
            conta_paga_lista.append(match_row["Conta_Norm"])
            data_paga_lista.append(match_row["Data_Norm"])

            if dias_diff == 0:
                obs_lista.append("Match exato de Valor e Data")
            elif dias_diff < 0:
                obs_lista.append(
                    f"Pago {abs(dias_diff)} dias antes do vencimento"
                )
            else:
                obs_lista.append(f"Pago {dias_diff} dias após o vencimento")
        else:
            status_lista.append("PENDENTE")
            obs_lista.append("")
            conta_paga_lista.append("-")
            data_paga_lista.append(None)

    df1["Status_Conciliacao"] = status_lista
    df1["Observacao"] = obs_lista
    df1["Conta_Paga"] = conta_paga_lista
    df1["Data_Paga"] = data_paga_lista

    for idx1, row1 in df1.iterrows():
        if df1.loc[idx1, "Status_Conciliacao"] != "PENDENTE":
            continue

        dt1 = row1["Data_Norm"]
        vl1 = row1["Valor_Norm"]
        margem = max(min(vl1 * 0.05, 50.0), 5.0)

        dt_min = dt1 - timedelta(days=35)
        dt_max = dt1 + timedelta(days=30)

        candidatos_p2 = df2[
            (~df2["_utilizado"])
            & (df2["Data_Norm"] >= dt_min)
            & (df2["Data_Norm"] <= dt_max)
            & (df2["Valor_Norm"] >= vl1 - margem)
            & (df2["Valor_Norm"] <= vl1 + margem)
        ].copy()

        if not candidatos_p2.empty:
            candidatos_p2["_diff"] = (candidatos_p2["Valor_Norm"] - vl1).abs()
            candidatos_p2 = candidatos_p2.sort_values(
                by=["_diff"], ascending=True
            )

            idx_m = candidatos_p2.index[0]
            match_row = candidatos_p2.loc[idx_m]
            df2.loc[idx_m, "_utilizado"] = True

            diff_val = match_row["Valor_Norm"] - vl1
            dias_diff = (match_row["Data_Norm"] - dt1).days

            df1.loc[idx1, "Status_Conciliacao"] = "COM_DIVERGENCIA"
            df1.loc[idx1, "Conta_Paga"] = match_row["Conta_Norm"]
            df1.loc[idx1, "Data_Paga"] = match_row["Data_Norm"]

            if diff_val > 0:
                obs = f"Provável Multa/Juros: +R$ {diff_val:.2f} ({abs(dias_diff)} dias de variação)"
            else:
                obs = f"Provável Desconto: -R$ {abs(diff_val):.2f} ({abs(dias_diff)} dias de variação)"

            df1.loc[idx1, "Observacao"] = obs
        else:
            df1.loc[idx1, "Status_Conciliacao"] = "NAO_ENCONTRADO"
            df1.loc[idx1, "Conta_Paga"] = "Pendente"
            df1.loc[idx1, "Data_Paga"] = None
            df1.loc[idx1, "Observacao"] = (
                "Nenhum correspondente no extrato nesta janela"
            )

    df_exatos = df1[df1["Status_Conciliacao"] == "EXATO"].copy()
    df_divergentes = df1[
        df1["Status_Conciliacao"] == "COM_DIVERGENCIA"
    ].copy()
    df_nao_encontrados = df1[
        df1["Status_Conciliacao"] == "NAO_ENCONTRADO"
    ].copy()

    wb = openpyxl.Workbook()

    cor_azul = "1A365D"
    cor_verde = "2F855A"
    cor_amarelo = "D69E2E"
    cor_vermelho = "C53030"
    cinza_fundo = "EDF2F7"
    cinza_borda = "CBD5E0"

    fonte_bold = Font(name="Segoe UI", size=10, bold=True)
    fonte_base = Font(name="Segoe UI", size=10)
    borda_fina = Border(
        left=Side(border_style="thin", color=cinza_borda),
        right=Side(border_style="thin", color=cinza_borda),
        top=Side(border_style="thin", color=cinza_borda),
        bottom=Side(border_style="thin", color=cinza_borda),
    )

    ws_resumo = wb.active
    ws_resumo.title = "Resumo da Conciliation"
    ws_resumo.views.sheetView[0].showGridLines = True

    ws_resumo.merge_cells("A1:E1")
    b = ws_resumo["A1"]
    b.value = (
        "PAINEL DE AUDITORIA - CONCILIAÇÃO PROVISIONADO vs EXTRATO BANCÁRIO"
    )
    b.font = Font(name="Segoe UI", size=12, bold=True, color="FFFFFF")
    b.fill = PatternFill(
        start_color=cor_azul, end_color=cor_azul, fill_type="solid"
    )
    b.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_resumo.row_dimensions[1].height = 28

    headers_resumo = [
        "STATUS",
        "QTD REGISTROS",
        "VALOR TOTAL (R$)",
        "% REGISTROS",
    ]
    ws_resumo.row_dimensions[3].height = 20
    for c_i, h in enumerate(headers_resumo, start=1):
        cell = ws_resumo.cell(row=3, column=c_i, value=h)
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            start_color="2B6CB0", end_color="2B6CB0", fill_type="solid"
        )
        cell.alignment = Alignment(horizontal="center", vertical="center")

    tot_qtd = len(df1)
    linhas_resumo = [
        (
            "MATCH EXATO",
            len(df_exatos),
            float(df_exatos["Valor_Norm"].sum()) if not df_exatos.empty else 0,
            cor_verde,
        ),
        (
            "MATCH COM DIVERGÊNCIA",
            len(df_divergentes),
            float(df_divergentes["Valor_Norm"].sum())
            if not df_divergentes.empty
            else 0,
            cor_amarelo,
        ),
        (
            "NÃO ENCONTRADOS / PENDENTES",
            len(df_nao_encontrados),
            float(df_nao_encontrados["Valor_Norm"].sum())
            if not df_nao_encontrados.empty
            else 0,
            cor_vermelho,
        ),
    ]

    for idx, (st, q, v, cor) in enumerate(linhas_resumo, start=4):
        ws_resumo.row_dimensions[idx].height = 20
        c1 = ws_resumo.cell(row=idx, column=1, value=st)
        c1.font = Font(name="Segoe UI", size=10, bold=True, color=cor)

        c2 = ws_resumo.cell(row=idx, column=2, value=q)
        c2.font = fonte_base
        c2.alignment = Alignment(horizontal="center")

        c3 = ws_resumo.cell(row=idx, column=3, value=v)
        c3.font = fonte_base
        c3.number_format = "R$ #,##0.00;[Red]-R$ #,##0.00"

        pct = (q / tot_qtd) if tot_qtd > 0 else 0
        c4 = ws_resumo.cell(row=idx, column=4, value=pct)
        c4.font = fonte_base
        c4.number_format = "0.0%"
        c4.alignment = Alignment(horizontal="center")

        for c in [c1, c2, c3, c4]:
            c.border = borda_fina

    r_tot = 7
    ws_resumo.row_dimensions[r_tot].height = 22
    ws_resumo.cell(row=r_tot, column=1, value="TOTAL PROVISIONADO").font = (
        fonte_bold
    )
    ws_resumo.cell(row=r_tot, column=2, value="=SUM(B4:B6)").font = fonte_bold
    ws_resumo.cell(row=r_tot, column=3, value="=SUM(C4:C6)").font = fonte_bold
    ws_resumo.cell(row=r_tot, column=3).number_format = (
        "R$ #,##0.00;[Red]-R$ #,##0.00"
    )
    ws_resumo.cell(row=r_tot, column=4, value="=SUM(D4:D6)").font = fonte_bold
    ws_resumo.cell(row=r_tot, column=4).number_format = "0.0%"

    for col_i in range(1, 5):
        c = ws_resumo.cell(row=r_tot, column=col_i)
        c.fill = PatternFill(
            start_color=cinza_fundo, end_color=cinza_fundo, fill_type="solid"
        )
        c.border = Border(
            top=Side(border_style="thin", color=cinza_borda),
            bottom=Side(border_style="double", color="1A202C"),
        )

    def _criar_aba_detalhe(titulo, df_dados, cor_banner):
        ws = wb.create_sheet(title=titulo)
        ws.views.sheetView[0].showGridLines = True

        headers = [
            "Conta Bancária",
            "Data Vencimento",
            "Data Pagamento (Efetiva/Provável)",
            "Fornecedor / Descrição",
            "Valor (R$)",
            "Status",
            "Análise da Baixa / Observação",
        ]

        ws.merge_cells("A1:G1")
        banner = ws["A1"]
        banner.value = f"DETALHAMENTO - {titulo.upper()}"
        banner.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        banner.fill = PatternFill(
            start_color=cor_banner, end_color=cor_banner, fill_type="solid"
        )
        banner.alignment = Alignment(
            horizontal="left", vertical="center", indent=1
        )
        ws.row_dimensions[1].height = 24

        ws.row_dimensions[3].height = 20
        for c_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=3, column=c_idx, value=h)
            cell.font = Font(
                name="Segoe UI", size=10, bold=True, color="FFFFFF"
            )
            cell.fill = PatternFill(
                start_color="4A5568", end_color="4A5568", fill_type="solid"
            )
            cell.alignment = Alignment(
                horizontal="left" if c_idx in [4, 7] else "center",
                vertical="center",
            )

        for i, row in df_dados.reset_index(drop=True).iterrows():
            l_atual = 4 + i
            ws.row_dimensions[l_atual].height = 19
            bg_zebra = PatternFill(
                start_color="F8FAFC" if i % 2 == 1 else "FFFFFF",
                fill_type="solid",
            )

            c1 = ws.cell(row=l_atual, column=1, value=str(row["Conta_Paga"]))
            c2 = ws.cell(row=l_atual, column=2, value=row["Data_Norm"])
            c3 = ws.cell(
                row=l_atual,
                column=3,
                value=row["Data_Paga"] if row["Data_Paga"] else "-",
            )
            col_forn_orig = df1.attrs.get("col_forn", df1.columns[2])
            c4 = ws.cell(
                row=l_atual,
                column=4,
                value=str(row[col_forn_orig])
                if pd.notna(row[col_forn_orig])
                else "-",
            )
            c5 = ws.cell(row=l_atual, column=5, value=_limpar_valor_monetario(row["Valor_Norm"]))
            c6 = ws.cell(row=l_atual, column=6, value=row["Status_Conciliacao"])
            c7 = ws.cell(row=l_atual, column=7, value=row["Observacao"])

            c1.alignment = Alignment(horizontal="center")
            c1.number_format = "@"
            c2.number_format = "DD/MM/YYYY"
            c2.alignment = Alignment(horizontal="center")

            if row["Data_Paga"]:
                c3.number_format = "DD/MM/YYYY"
            c3.alignment = Alignment(horizontal="center")

            c4.alignment = Alignment(horizontal="left")
            c5.number_format = "R$ #,##0.00;[Red]-R$ #,##0.00"
            c5.alignment = Alignment(horizontal="right")
            c6.alignment = Alignment(horizontal="center")
            c7.alignment = Alignment(horizontal="left")

            for c in [c1, c2, c3, c4, c5, c6, c7]:
                c.border = borda_fina
                c.font = fonte_base
                c.fill = bg_zebra

        if not df_dados.empty:
            l_tot = 4 + len(df_dados)
            ws.row_dimensions[l_tot].height = 22
            ws.cell(row=l_tot, column=1, value="TOTAL").font = fonte_bold
            c_sum = ws.cell(row=l_tot, column=5, value=f"=SUM(E4:E{l_tot-1})")
            c_sum.font = fonte_bold
            c_sum.number_format = "R$ #,##0.00;[Red]-R$ #,##0.00"

            for col_i in range(1, 8):
                c = ws.cell(row=l_tot, column=col_i)
                c.fill = PatternFill(
                    start_color=cinza_fundo,
                    end_color=cinza_fundo,
                    fill_type="solid",
                )
                c.border = Border(
                    top=Side(border_style="thin", color=cinza_borda),
                    bottom=Side(border_style="double", color="1A202C"),
                )

        for col_idx in range(1, 8):
            comp_max = max(
                len(str(ws.cell(row=r, column=col_idx).value or ""))
                for r in range(3, ws.max_row + 1)
            )
            ws.column_dimensions[get_column_letter(col_idx)].width = max(
                comp_max + 4, 15
            )

    _criar_aba_detalhe("Match Exato", df_exatos, cor_verde)
    _criar_aba_detalhe("Com Divergência", df_divergentes, cor_amarelo)
    _criar_aba_detalhe("Não Encontrados", df_nao_encontrados, cor_vermelho)

    for col_idx in range(1, 5):
        ws_resumo.column_dimensions[get_column_letter(col_idx)].width = 35

    output_stream = io.BytesIO()
    wb.save(output_stream)
    output_stream.seek(0)
    return output_stream


# ==============================================================================
# VIEW PRINCIPAL DO DJANGO
# ==============================================================================
def comparar_planilhas_view(request):
    if request.method == "POST":
        arq1 = request.FILES.get("planilha_1")
        arq2 = request.FILES.get("planilha_2")
        modo = request.POST.get("modo_auditoria", "BANCO")

        if not arq1 or not arq2:
            messages.error(
                request, "É necessário enviar as duas planilhas (.xlsx)."
            )
            return render(request, "contas_pagar/comparar_planilhas.html")

        if not arq1.name.lower().endswith(
            ".xlsx"
        ) or not arq2.name.lower().endswith(".xlsx"):
            messages.error(
                request, "Apenas arquivos no formato .xlsx são suportados."
            )
            return render(request, "contas_pagar/comparar_planilhas.html")

        try:
            if modo == "BOLETOS":
                excel_bytes = conciliar_boletos_em_maos(arq1, arq2)
                nome_download = f"Boletos_Faltantes_Importar_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            else:
                excel_bytes = conciliar_extrato_banco(arq1, arq2)
                nome_download = f"Conciliacao_Bradesco_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

            response = HttpResponse(
                excel_bytes.getvalue(),
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            response["Content-Disposition"] = (
                f'attachment; filename="{nome_download}"'
            )
            return response

        except Exception as e:
            messages.error(
                request, f"Erro ao processar e comparar as planilhas: {str(e)}"
            )

    return render(request, "contas_pagar/comparar_planilhas.html")