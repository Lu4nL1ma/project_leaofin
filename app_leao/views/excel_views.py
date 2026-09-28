from datetime import datetime
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.http import HttpResponse, JsonResponse
from django.db import transaction
from django.views.decorators.http import require_POST
from app_leao.models import ContaPagar, Fornecedor, Categoria, BancoSaldo
from app_leao.utils.formatters import extrair_texto, limpar_cnpj, formatar_cnpj, parse_valor, parse_data

@require_POST
def importar_xlsx(request):
    excel_file = request.FILES.get('arquivo_xlsx') or request.FILES.get('arquivo_excel')
    if not excel_file:
        return JsonResponse({'sucesso': False, 'erro': 'Nenhum arquivo enviado.'}, status=400)

    try:
        wb = openpyxl.load_workbook(excel_file, data_only=True)
        sheet = wb.active
        header = [extrair_texto(cell.value).lower() for cell in sheet[1]]
        
        def achar_coluna(nomes):
            for nome in nomes:
                if nome in header:
                    return header.index(nome)
            return None

        idx_cnpj = achar_coluna(['cnpj', 'cnpj fornecedor', 'cpf/cnpj', 'fornecedor_cnpj'])
        idx_nf = achar_coluna(['nota fiscal', 'nf', 'numero nf', 'num_nota', 'nota_fiscal'])
        idx_linha = achar_coluna(['linha digitavel', 'boleto', 'codigo de barras', 'linha_digitavel'])
        idx_valor = achar_coluna(['valor', 'valor (r$)', 'valor total', 'valor_total'])
        idx_venc = achar_coluna(['vencimento', 'data vencimento', 'dt vencimento', 'data_vencimento', 'venc'])
        idx_categoria = achar_coluna(['categoria', 'cat', 'categoria_nome'])
        idx_banco = achar_coluna(['banco', 'conta bancaria', 'banco_nome', 'bancosaldo'])
        idx_parcela = achar_coluna(['parcela', 'nº parcela', 'numero parcela', 'parcela_numero'])

        colunas_faltantes = []
        if idx_cnpj is None: colunas_faltantes.append('CNPJ')
        if idx_valor is None: colunas_faltantes.append('Valor')
        if idx_venc is None: colunas_faltantes.append('Vencimento')
        if idx_categoria is None: colunas_faltantes.append('Categoria')
        if idx_banco is None: colunas_faltantes.append('Banco')
        if idx_parcela is None: colunas_faltantes.append('Parcela')
        if colunas_faltantes:
            return JsonResponse({'sucesso': False, 'erro': f'Planilha inválida. Faltam: {", ".join(colunas_faltantes)}.'}, status=400)

        contas_novas = []
        contas_atualizadas = 0
        erros = []

        with transaction.atomic():
            for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                if not any(row): 
                    continue

                cnpj_raw = row[idx_cnpj] if idx_cnpj < len(row) else None
                cnpj_limpo = limpar_cnpj(cnpj_raw)
                if not cnpj_limpo:
                    erros.append(f"Linha {row_idx}: CNPJ ausente ou inválido.")
                    continue

                cnpj_formatado = formatar_cnpj(cnpj_limpo)
                fornecedor = Fornecedor.objects.filter(cnpj__in=[cnpj_limpo, cnpj_formatado]).first()
                if not fornecedor:
                    fornecedor = Fornecedor.objects.filter(cnpj__icontains=cnpj_limpo).first()
                if not fornecedor:
                    erros.append(f"Linha {row_idx}: Fornecedor com CNPJ {cnpj_limpo} não encontrado.")
                    continue

                nome_categoria = extrair_texto(row[idx_categoria] if idx_categoria < len(row) else "")
                categoria = Categoria.objects.filter(nome__iexact=nome_categoria).first() if nome_categoria else None
                if not categoria:
                    erros.append(f"Linha {row_idx}: Categoria '{nome_categoria}' não encontrada.")
                    continue

                nome_banco = extrair_texto(row[idx_banco] if idx_banco < len(row) else "")
                banco_saldo = None
                if hasattr(BancoSaldo, 'nome'):
                    banco_saldo = BancoSaldo.objects.filter(nome__icontains=nome_banco).first()
                if not banco_saldo and hasattr(BancoSaldo, 'banco'):
                    banco_saldo = BancoSaldo.objects.filter(banco__icontains=nome_banco).first()
                if not banco_saldo and nome_banco.isdigit():
                    banco_saldo = BancoSaldo.objects.filter(pk=nome_banco).first()
                if not banco_saldo:
                    erros.append(f"Linha {row_idx}: Banco/Conta '{nome_banco}' não encontrado.")
                    continue

                valor = parse_valor(row[idx_valor] if idx_valor < len(row) else 0)
                vencimento = parse_data(row[idx_venc] if idx_venc < len(row) else None)
                if not vencimento:
                    erros.append(f"Linha {row_idx}: Data de vencimento inválida.")
                    continue

                nf = extrair_texto(row[idx_nf]) if idx_nf is not None and idx_nf < len(row) else ""
                linha_dig = extrair_texto(row[idx_linha]) if idx_linha is not None and idx_linha < len(row) else ""
                parcela = extrair_texto(row[idx_parcela]) if idx_parcela is not None and idx_parcela < len(row) else "01/01"

                conta_existente = ContaPagar.objects.filter(fornecedor=fornecedor, vencimento=vencimento, valor=valor).first()
                if conta_existente:
                    atualizou = False
                    if not conta_existente.nota_fiscal and nf:
                        conta_existente.nota_fiscal = nf
                        atualizou = True
                    if not conta_existente.linha_digitavel and linha_dig:
                        conta_existente.linha_digitavel = linha_dig
                        atualizou = True
                    if atualizou:
                        conta_existente.save()
                        contas_atualizadas += 1
                    continue

                contas_novas.append(ContaPagar(
                    fornecedor=fornecedor,
                    categoria=categoria,
                    banco=banco_saldo,
                    nota_fiscal=nf,
                    linha_digitavel=linha_dig,
                    valor=valor,
                    vencimento=vencimento,
                    status='Pendente',
                    parcela=parcela
                ))

            if contas_novas:
                ContaPagar.objects.bulk_create(contas_novas)

        return JsonResponse({'sucesso': True, 'importados': len(contas_novas), 'atualizados': contas_atualizadas, 'erros': erros})
    except Exception as e:
        return JsonResponse({'sucesso': False, 'erro': f'Erro ao ler arquivo: {str(e)}'}, status=500)

def baixar_planilha_padrao(request):
    workbook = Workbook()
    aba = workbook.active
    aba.title = "Modelo"

    cabecalho = ["vencimento", "fornecedor", "categoria", "banco", "parcela", "valor", "observação", "status", "nota_fiscal", "linha_digitavel"]
    aba.append(cabecalho)
    aba.append(["12/12/2026", "Seu Fornecedor", "Energia", "Nome Banco", "1/1", 150.00, "Sua Observação", "Pendente","0000", "00000000000000000000000000000"])

    for coluna in aba.columns:
        maior_largura = max(len(str(celula.value)) for celula in coluna)
        aba.column_dimensions[coluna[0].column_letter].width = maior_largura + 4

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="modelo_contas_a_pagar.xlsx"'
    workbook.save(response)
    return response

def exportar_contas_pagar_excel(request):
    tipo_exportacao = request.GET.get('tipo', 'tudo')
    data_inicio = request.GET.get('data_inicio')
    data_fim = request.GET.get('data_fim')

    queryset = ContaPagar.objects.all().order_by('vencimento')
    if tipo_exportacao == 'periodo' and data_inicio and data_fim:
        queryset = queryset.filter(vencimento__range=[data_inicio, data_fim])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Contas a Pagar"
    ws.views.sheetView[0].showGridLines = True

    font_header = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    fill_header = PatternFill(start_color='1F497D', end_color='1F497D', fill_type='solid')
    font_dados = Font(name='Segoe UI', size=10)
    fill_zebra = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    fill_white = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
    font_total = Font(name='Segoe UI', size=11, bold=True, color='1F497D')
    fill_total = PatternFill(start_color='DCE6F1', end_color='DCE6F1', fill_type='solid')

    border_thin = Border(
        left=Side(style='thin', color='E2E8F0'), right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'), bottom=Side(style='thin', color='E2E8F0')
    )
    border_total = Border(top=Side(style='thin', color='1F497D'), bottom=Side(style='double', color='1F497D'))

    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')

    headers = ['ID', 'Status', 'Vencimento', 'Fornecedor', 'Categoria', 'Nota Fiscal', 'Parcela', 'Valor (R$)', 'Banco', 'Linha Digitável', 'Observação']
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_thin

    start_row = 2
    for idx, conta in enumerate(queryset, start=start_row):
        is_even = (idx - start_row) % 2 == 0
        current_fill = fill_white if is_even else fill_zebra

        venc = getattr(conta, 'vencimento', None)
        venc_str = venc.strftime('%d/%m/%Y') if venc else '-'
        valor = float(getattr(conta, 'valor', 0) or 0)

        row_data = [
            conta.id, str(getattr(conta, 'status', '-')).upper(), venc_str,
            str(getattr(conta, 'fornecedor', '-')), str(getattr(conta, 'categoria', '-')),
            str(getattr(conta, 'nota_fiscal', '-')), str(getattr(conta, 'parcela', '-')),
            valor, str(getattr(conta, 'banco', '-')), str(getattr(conta, 'linha_digitavel', '-')),
            str(getattr(conta, 'observacao', '-'))
        ]
        ws.append(row_data)

        for col_num in range(1, len(headers) + 1):
            c = ws.cell(row=idx, column=col_num)
            c.font = font_dados
            c.fill = current_fill
            c.border = border_thin
            if col_num in (1, 2, 3, 6, 7):
                c.alignment = align_center
            elif col_num in (4, 5, 9, 10, 11):
                c.alignment = align_left
            elif col_num == 8:
                c.alignment = align_right
                c.number_format = 'R$ #,##0.00'

    last_row = start_row + len(queryset) - 1 if len(queryset) > 0 else start_row
    if len(queryset) > 0:
        total_row = last_row + 1
        ws.cell(row=total_row, column=1, value="TOTAL")
        ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=7)
        ws.cell(row=total_row, column=1).alignment = Alignment(horizontal='right', vertical='center')

        sum_cell = ws.cell(row=total_row, column=8, value=f"=SUM(H{start_row}:H{last_row})")
        sum_cell.number_format = 'R$ #,##0.00'
        sum_cell.alignment = align_right

        for col_num in range(1, len(headers) + 1):
            c = ws.cell(row=total_row, column=col_num)
            c.font = font_total
            c.fill = fill_total
            c.border = border_total

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    ws.freeze_panes = 'A2'
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    filename = f"contas_a_pagar_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response