import io
import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal
from typing import List

from ofxparse import OfxParser
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd


# ==============================================================================
# DICIONÁRIO / DEPARA DE CONTAS CORRENTES
# ==============================================================================
MAPEAMENTO_CONTA_CORRENTE = {
    "40147": "Bradesco (Matriz)",
    "40436": "Bradesco (Aeroporto)",
    "41835": "Bradesco (Castanheira)",
    "41828": "Bradesco (Baenão)",
    "41818": "Bradesco (Pátio Belém)",
}


def _converter_nome_conta(valor_bruto) -> str:
    """Converte códigos numéricos de conta do OFX para o nome formatado da agência/loja."""
    if not valor_bruto or pd.isna(valor_bruto):
        return "Não Identificado"

    val_str = str(valor_bruto).strip().replace(".0", "")

    for cod, nome in MAPEAMENTO_CONTA_CORRENTE.items():
        if cod in val_str:
            return nome

    return val_str


def _remover_acentos(texto: str) -> str:
    """Remove caracteres acentuados para evitar falha no parser interno do ofxparse."""
    return (
        unicodedata.normalize("NFKD", texto)
        .encode("ascii", "ignore")
        .decode("ascii")
    )


def _carregar_buffer_ofx(arquivo_django) -> io.BytesIO:
    if hasattr(arquivo_django, "seek"):
        arquivo_django.seek(0)

    raw = arquivo_django.read()

    texto = None
    for enc in ("utf-8", "latin-1", "cp1252", "iso-8859-1"):
        try:
            texto = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue

    if texto is None:
        texto = raw.decode("latin-1", errors="replace")

    # Higieniza caracteres acentuados
    texto_limpo = _remover_acentos(texto)

    # Padroniza tags de cabeçalho para USASCII
    texto_limpo = re.sub(
        r"CHARSET:\s*(?:1252|NONE|UTF-8)",
        "CHARSET:USASCII",
        texto_limpo,
        flags=re.IGNORECASE,
    )
    texto_limpo = re.sub(
        r"ENCODING:\s*UTF-8", "ENCODING:USASCII", texto_limpo, flags=re.IGNORECASE
    )

    return io.BytesIO(texto_limpo.encode("ascii", "ignore"))


def extrair_dados_ofx(arquivo_django):
    buffer = _carregar_buffer_ofx(arquivo_django)
    ofx = OfxParser.parse(buffer)

    account = getattr(ofx, "account", None)
    if not account and hasattr(ofx, "accounts") and ofx.accounts:
        account = ofx.accounts[0]

    stmt = getattr(account, "statement", None)
    conta_bruta = (
        str(account.account_id) if account and account.account_id else "SEM_CONTA"
    )
    
    # Aplica a conversão para o nome formatado da conta
    conta_num = _converter_nome_conta(conta_bruta)

    banco_cod = (
        account.routing_number if account and account.routing_number else "0237"
    )

    d_ini = stmt.start_date.date() if (stmt and stmt.start_date) else None
    d_fim = stmt.end_date.date() if (stmt and stmt.end_date) else None
    saldo = float(stmt.balance) if (stmt and stmt.balance is not None) else 0.0

    transacoes = []
    if stmt and stmt.transactions:
        for tx in stmt.transactions:
            valor = (
                float(tx.amount)
                if isinstance(tx.amount, (Decimal, float, int))
                else 0.0
            )
            d_tx = (
                tx.date.date()
                if isinstance(tx.date, datetime)
                else (tx.date if isinstance(tx.date, date) else None)
            )

            transacoes.append(
                {
                    "Conta Corrente": conta_num,
                    "Data": d_tx,
                    "Tipo": (
                        "Crédito" if (tx.type == "credit" or valor > 0) else "Débito"
                    ),
                    "Descrição / Histórico": (
                        tx.memo.strip() if tx.memo else "SEM HISTÓRICO"
                    ),
                    "Nº Documento": str(tx.checknum or "-"),
                    "ID Lançamento": str(tx.id or "-"),
                    "Valor (R$)": valor,
                }
            )

    meta = {
        "conta": conta_num,
        "banco": banco_cod,
        "data_inicio": d_ini,
        "data_fim": d_fim,
        "saldo_final": saldo,
    }
    return transacoes, meta


def gerar_excel_consolidado_em_memoria(lista_arquivos: List) -> io.BytesIO:
    todas_transacoes = []
    contas_metadados = []

    for arq in lista_arquivos:
        txs, meta = extrair_dados_ofx(arq)
        todas_transacoes.extend(txs)
        contas_metadados.append(meta)

    df = pd.DataFrame(todas_transacoes)
    if not df.empty:
        df = df.sort_values(
            by=["Data", "Conta Corrente"], ascending=[True, True]
        ).reset_index(drop=True)

    datas_inicio = [m["data_inicio"] for m in contas_metadados if m["data_inicio"]]
    datas_fim = [m["data_fim"] for m in contas_metadados if m["data_fim"]]
    periodo_geral_ini = (
        min(datas_inicio).strftime("%d/%m/%Y") if datas_inicio else "N/D"
    )
    periodo_geral_fim = (
        max(datas_fim).strftime("%d/%m/%Y") if datas_fim else "N/D"
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Consolidado OFX"
    ws.views.sheetView[0].showGridLines = True

    cor_primaria = "1A365D"
    cor_cabecalho = "2B6CB0"
    cinza_fundo = "EDF2F7"
    cinza_borda = "CBD5E0"

    header_fill = PatternFill(
        start_color=cor_cabecalho, end_color=cor_cabecalho, fill_type="solid"
    )
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    fonte_bold = Font(name="Segoe UI", size=10, bold=True)
    fonte_base = Font(name="Segoe UI", size=10)
    fonte_credito = Font(name="Segoe UI", size=10, color="2B6CB0", bold=True)
    fonte_debito = Font(name="Segoe UI", size=10, color="C53030", bold=True)

    borda_fina = Border(
        left=Side(border_style="thin", color=cinza_borda),
        right=Side(border_style="thin", color=cinza_borda),
        top=Side(border_style="thin", color=cinza_borda),
        bottom=Side(border_style="thin", color=cinza_borda),
    )

    # Banner
    ws.merge_cells("A1:G1")
    banner = ws["A1"]
    banner.value = "RELATÓRIO CONSOLIDADO DE CONTAS - MULTI-EXTRATO (5 CONTAS)"
    banner.font = Font(name="Segoe UI", size=12, bold=True, color="FFFFFF")
    banner.fill = PatternFill(
        start_color=cor_primaria, end_color=cor_primaria, fill_type="solid"
    )
    banner.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 26

    # Metadados
    ws.row_dimensions[2].height = 20
    ws.cell(row=2, column=1, value="Período Geral:").font = fonte_bold
    ws.cell(
        row=2, column=2, value=f"{periodo_geral_ini} até {periodo_geral_fim}"
    ).font = fonte_base
    ws.cell(row=2, column=4, value="Contas Processadas:").font = fonte_bold
    ws.cell(row=2, column=5, value=len(contas_metadados)).font = fonte_bold

    # Painel dos Saldos Finais por Conta
    linha_painel = 4
    ws.row_dimensions[linha_painel].height = 20
    ws.cell(row=linha_painel, column=1, value="CONTA").font = fonte_bold
    ws.cell(row=linha_painel, column=2, value="SALDO FINAL").font = fonte_bold
    ws.cell(row=linha_painel, column=1).fill = PatternFill(
        start_color=cinza_fundo, end_color=cinza_fundo, fill_type="solid"
    )
    ws.cell(row=linha_painel, column=2).fill = PatternFill(
        start_color=cinza_fundo, end_color=cinza_fundo, fill_type="solid"
    )

    for idx, c_meta in enumerate(contas_metadados):
        r = linha_painel + 1 + idx
        ws.row_dimensions[r].height = 18
        c_conta = ws.cell(row=r, column=1, value=str(c_meta["conta"]))
        c_conta.font = fonte_base
        c_conta.alignment = Alignment(horizontal="left", vertical="center")

        c_saldo = ws.cell(row=r, column=2, value=c_meta["saldo_final"])
        c_saldo.font = fonte_base
        c_saldo.number_format = "R$ #,##0.00;[Red]-R$ #,##0.00"
        c_saldo.alignment = Alignment(horizontal="right", vertical="center")

    # Linha com Soma dos Saldos
    linha_saldo_tot = linha_painel + len(contas_metadados) + 1
    ws.row_dimensions[linha_saldo_tot].height = 20
    c_tot_txt = ws.cell(row=linha_saldo_tot, column=1, value="SOMA DOS SALDOS:")
    c_tot_txt.font = fonte_bold

    c_tot_val = ws.cell(
        row=linha_saldo_tot,
        column=2,
        value=f"=SUM(B{linha_painel+1}:B{linha_saldo_tot-1})",
    )
    c_tot_val.font = fonte_bold
    c_tot_val.number_format = "R$ #,##0.00;[Red]-R$ #,##0.00"
    c_tot_val.alignment = Alignment(horizontal="right", vertical="center")

    # Tabela com Transações
    colunas_tabela = [
        "Conta Corrente",
        "Data",
        "Tipo",
        "Descrição / Histórico",
        "Nº Documento",
        "ID Lançamento",
        "Valor (R$)",
    ]
    linha_cabecalho = linha_saldo_tot + 2
    ws.row_dimensions[linha_cabecalho].height = 22

    for c_idx, nome in enumerate(colunas_tabela, start=1):
        c = ws.cell(row=linha_cabecalho, column=c_idx, value=nome)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(
            horizontal="left" if c_idx in [1, 4] else "center", vertical="center"
        )

    for i, row in df.iterrows():
        l_atual = linha_cabecalho + 1 + i
        ws.row_dimensions[l_atual].height = 19
        bg_zebra = PatternFill(
            start_color="F8FAFC" if i % 2 == 1 else "FFFFFF", fill_type="solid"
        )

        c_conta = ws.cell(row=l_atual, column=1, value=str(row["Conta Corrente"]))
        c_data = ws.cell(row=l_atual, column=2, value=row["Data"])
        c_tipo = ws.cell(row=l_atual, column=3, value=row["Tipo"])
        c_desc = ws.cell(row=l_atual, column=4, value=row["Descrição / Histórico"])
        c_doc = ws.cell(row=l_atual, column=5, value=row["Nº Documento"])
        c_id = ws.cell(row=l_atual, column=6, value=row["ID Lançamento"])
        c_val = ws.cell(row=l_atual, column=7, value=row["Valor (R$)"])

        c_conta.alignment = Alignment(horizontal="left", vertical="center")

        c_data.number_format = "DD/MM/YYYY"
        c_data.alignment = Alignment(horizontal="center", vertical="center")

        c_tipo.alignment = Alignment(horizontal="center", vertical="center")
        c_desc.alignment = Alignment(horizontal="left", vertical="center")
        c_doc.alignment = Alignment(horizontal="center", vertical="center")
        c_id.alignment = Alignment(horizontal="center", vertical="center")

        c_val.number_format = "R$ #,##0.00;[Red]-R$ #,##0.00"
        c_val.alignment = Alignment(horizontal="right", vertical="center")
        c_val.font = fonte_credito if row["Valor (R$)"] >= 0 else fonte_debito

        for c in [c_conta, c_data, c_tipo, c_desc, c_doc, c_id, c_val]:
            c.border = borda_fina
            if c != c_val:
                c.font = fonte_base
            c.fill = bg_zebra

    # Rodapé Total das Movimentações
    if not df.empty:
        l_total = linha_cabecalho + len(df) + 1
        ws.row_dimensions[l_total].height = 22
        ws.cell(
            row=l_total, column=1, value="TOTAL DAS MOVIMENTAÇÕES"
        ).font = fonte_bold

        tot_mov = ws.cell(
            row=l_total,
            column=7,
            value=f"=SUM(G{linha_cabecalho+1}:G{l_total-1})",
        )
        tot_mov.font = fonte_bold
        tot_mov.number_format = "R$ #,##0.00;[Red]-R$ #,##0.00"
        tot_mov.alignment = Alignment(horizontal="right", vertical="center")

        borda_dupla = Border(
            top=Side(border_style="thin", color=cinza_borda),
            bottom=Side(border_style="double", color="1A202C"),
        )
        for col_idx in range(1, 8):
            ws.cell(row=l_total, column=col_idx).border = borda_dupla
            ws.cell(row=l_total, column=col_idx).fill = PatternFill(
                start_color=cinza_fundo, end_color=cinza_fundo, fill_type="solid"
            )

    # Ajuste de largura das colunas
    for col_idx in range(1, 8):
        comp_max = 0
        for r_idx in range(linha_cabecalho, ws.max_row + 1):
            val = ws.cell(row=r_idx, column=col_idx).value
            if val is not None:
                if isinstance(val, (datetime, date)):
                    comp_max = max(comp_max, 10)
                elif isinstance(val, float):
                    comp_max = max(comp_max, len(f"R$ {val:,.2f}"))
                elif str(val).startswith("="):
                    comp_max = max(comp_max, 14)
                else:
                    comp_max = max(comp_max, len(str(val)))

        ws.column_dimensions[get_column_letter(col_idx)].width = max(
            comp_max + 4, 15
        )

    output_stream = io.BytesIO()
    wb.save(output_stream)
    output_stream.seek(0)
    return output_stream